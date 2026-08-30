"""The seed-review gate: what may reach `eval/gold/`, and what may not.

These tests guard one property above all others, and it is the property that
would fail silently if it broke: **an unapproved row must never reach the gold
set.** A seed pack's rows validate, resolve and hash correctly whether or not a
person read them, so no other check in this repo would notice the regression.
The gold set would simply grow, and every later stage would be measured against
machine writing.

The rest is the vocabulary around that gate — verdict normalisation that must
happen, verdict repair that must not, and a date format that must not be
guessed.
"""

from __future__ import annotations

import pytest

from tm_knowledge.provenance import ReviewStatus
from tm_knowledge.stage0 import review


def make_review(verdict: str = "correct", correction: str | None = None) -> review.Review:
    return review.Review(
        seed_id="SEED-GC-001",
        why_this_example="why",
        passage="passage",
        verdict=verdict,
        correction=correction,
    )


# ---------------------------------------------------------------------------
# The gate
# ---------------------------------------------------------------------------


def test_correct_and_signed_is_the_only_way_into_gold():
    destination, status, _, flag = review.route(make_review(), "TC", "2026-08-25")
    assert destination == "gold"
    assert status is ReviewStatus.APPROVED
    assert flag is None


@pytest.mark.parametrize(
    "verdict, approved_by, approved_date",
    [
        ("correct", None, None),        # reviewed, never signed
        ("correct", "TC", None),        # half-signed
        ("correct", None, "2026-08-25"),
        ("amend", "TC", "2026-08-25"),  # signed, but the expert wants it changed
        ("reject", "TC", "2026-08-25"),
        ("unreviewed", "TC", "2026-08-25"),
        ("unreviewed", None, None),
    ],
)
def test_everything_else_is_held(verdict, approved_by, approved_date):
    destination, status, reason, _ = review.route(
        make_review(verdict), approved_by, approved_date
    )
    assert destination == "review"
    assert status is not ReviewStatus.APPROVED
    assert reason, "a held row must always say why it was held"


def test_a_signature_alone_does_not_approve_an_unread_row():
    """The failure mode worth naming: a signed sheet with verdicts left blank.

    Someone could reasonably sign the bottom of a sheet meaning "I looked at
    this lot" without marking each row. That is not a verdict, and treating it
    as one would promote unread machine writing under a real person's name.
    """
    destination, status, _, _ = review.route(
        make_review("unreviewed"), "TC", "2026-08-25"
    )
    assert destination == "review"
    assert status is ReviewStatus.CANDIDATE


def test_correct_plus_a_correction_is_a_conflict_not_an_approval():
    destination, _, _, flag = review.route(
        make_review("correct", correction="duplicate"), "TC", "2026-08-25"
    )
    assert destination == "review"
    assert flag is not None, "the reviewer must be told the two cells disagree"


def test_rejection_keeps_its_status_rather_than_disappearing():
    _, status, _, _ = review.route(make_review("reject"), "TC", "2026-08-25")
    assert status is ReviewStatus.REJECTED


# ---------------------------------------------------------------------------
# Verdicts
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "written, expected",
    [
        ("correct", "correct"),
        ("Correct", "correct"),
        ("CORRECT", "correct"),
        ("amend ", "amend"),
        ("  Reject", "reject"),
        ("", "unreviewed"),
        ("   ", "unreviewed"),
        (None, "unreviewed"),
    ],
)
def test_case_and_whitespace_are_normalised(written, expected):
    """Neither carries judgement — it is one word typed into a spreadsheet."""
    assert review.read_verdict(written) == expected


@pytest.mark.parametrize("typo", ["corrrect", "rejext", "amended", "ok", "yes", "3"])
def test_a_verdict_that_is_not_a_verdict_is_refused_not_repaired(typo):
    """`corrrect` is obviously `correct` and is still not read as one.

    Guessing here would mean a verdict in `review/decisions/` is not necessarily
    a verdict a person gave, which is the one thing this directory promises.
    """
    with pytest.raises(review.UnreadableVerdict):
        review.read_verdict(typo)


def test_the_refusal_names_the_cell_and_the_options():
    with pytest.raises(review.UnreadableVerdict) as caught:
        review.read_verdict("corrrect")
    message = str(caught.value)
    assert "corrrect" in message
    assert "correct" in message and "reject" in message


# ---------------------------------------------------------------------------
# Dates
# ---------------------------------------------------------------------------


def test_australian_day_order_is_converted_when_the_reading_is_forced():
    assert review.read_approved_date("25/08/2026") == "2026-08-25"
    assert review.read_approved_date("20/08/2026") == "2026-08-20"


def test_iso_passes_through():
    assert review.read_approved_date("2026-08-25") == "2026-08-25"


def test_a_datetime_cell_is_taken_as_it_stands():
    from datetime import datetime

    assert review.read_approved_date(datetime(2026, 8, 25, 9, 30)) == "2026-08-25"


@pytest.mark.parametrize("ambiguous", ["05/08/2026", "01/02/2026", "12/11/2026"])
def test_an_ambiguous_date_is_refused_rather_than_read_as_australian(ambiguous):
    """Australian order is the convention and is still not assumed.

    An approval date is half of the approval artefact (ADR-0039), so a guessed
    one is a falsified audit trail rather than a formatting inconvenience.
    """
    with pytest.raises(review.AmbiguousDate):
        review.read_approved_date(ambiguous)


def test_blank_dates_stay_blank():
    assert review.read_approved_date(None) is None
    assert review.read_approved_date("  ") is None


def test_nonsense_is_refused():
    with pytest.raises(ValueError):
        review.read_approved_date("last Tuesday")


# ---------------------------------------------------------------------------
# The audit trail
# ---------------------------------------------------------------------------


def _decision(**overrides):
    base = dict(
        record_id="GE-0003",
        record_type="gold_entity",
        sheet="entities",
        row=4,
        verdict="reject",
        review_status=ReviewStatus.REJECTED,
        destination="review",
        reason="verdict is `reject`",
        correction="Better entity definition further down",
        approved_by=None,
        approved_date=None,
        seed_id="SEED-GE-0003",
        record={"id": "GE-0003", "surface": "this section"},
    )
    base.update(overrides)
    return review.Decision(**base)


def test_a_rejection_carries_the_whole_record_not_just_its_id():
    entry = _decision().as_dict()
    assert entry["record"]["surface"] == "this section"
    assert entry["correction"] == "Better entity definition further down"


def test_every_entry_is_labelled_as_machine_written():
    """Rule 8: model output is labelled wherever it is stored."""
    assert _decision().as_dict()["origin"] == "machine_seed"


def test_decisions_are_written_one_file_per_record_type(tmp_path):
    outcomes = review.write_decisions(
        [_decision(), _decision(record_id="GE-0004")],
        tmp_path,
        source="workbook.xlsx",
        write=True,
    )
    assert [outcome for _, outcome in outcomes] == ["written"]
    written = (tmp_path / "gold-entity.yaml").read_text(encoding="utf-8")
    assert "workbook.xlsx" in written
    assert "NOT approved knowledge" in written


def test_rewriting_unchanged_decisions_changes_nothing(tmp_path):
    review.write_decisions([_decision()], tmp_path, source="w.xlsx", write=True)
    outcomes = review.write_decisions([_decision()], tmp_path, source="w.xlsx", write=True)
    assert [outcome for _, outcome in outcomes] == ["unchanged"]


def test_a_dry_run_writes_nothing(tmp_path):
    outcomes = review.write_decisions([_decision()], tmp_path, source="w.xlsx")
    assert [outcome for _, outcome in outcomes] == ["would write"]
    assert not list(tmp_path.iterdir())


# ---------------------------------------------------------------------------
# End to end, through the transcriber
# ---------------------------------------------------------------------------

pytest.importorskip("openpyxl", reason="the intake path needs `pip install -e '.[intake]'`")

from tm_knowledge.config import REPO_ROOT  # noqa: E402
from tm_knowledge.stage0 import goldset, transcribe, workbook  # noqa: E402

SOUND = REPO_ROOT / "tests" / "fixtures" / "harness" / "sound"


def _seed_pack(tmp_path, verdicts, *, sign=True, corrections=None, name="seed.xlsx"):
    """A workbook of sound records, given review columns and a verdict each."""
    from openpyxl import load_workbook

    gold = goldset.load(SOUND)
    records = {rt: list(rs) for rt, rs in gold.records.items()}
    book = workbook.build(generated="2026-08-25")
    workbook.fill(book, records)
    path = tmp_path / name
    book.save(path)

    book = load_workbook(path)
    for sheet in book.worksheets:
        headers = [c.value for c in sheet[1]]
        if "id" not in headers and "parent_id" not in headers:
            continue
        start = sheet.max_column + 1
        for offset, header in enumerate(review.REVIEW_HEADERS):
            sheet.cell(row=1, column=start + offset, value=header)
        verdict_col = start + review.REVIEW_HEADERS.index("verdict")
        correction_col = start + review.REVIEW_HEADERS.index("correction")
        seed_col = start + review.REVIEW_HEADERS.index("seed_id")
        for row in range(2, sheet.max_row + 1):
            identifier = sheet.cell(row=row, column=1).value
            if identifier is None:
                continue
            sheet.cell(row=row, column=verdict_col, value=verdicts.get(identifier, "correct"))
            sheet.cell(row=row, column=seed_col, value=f"SEED-{identifier}")
            if corrections and identifier in corrections:
                sheet.cell(row=row, column=correction_col, value=corrections[identifier])
            if not sign and "approved_by" in headers:
                # `cell(value=None)` is a no-op in openpyxl — assign to clear.
                sheet.cell(row=row, column=headers.index("approved_by") + 1).value = None
                sheet.cell(row=row, column=headers.index("approved_date") + 1).value = None
    book.save(path)
    return path


def test_review_columns_are_accepted_where_an_unknown_column_is_not(tmp_path):
    """The five named columns read; anything else still stops the read."""
    path = _seed_pack(tmp_path, {})
    result = transcribe.read_workbook(path)
    assert result.decisions, "the review columns should have been read"


def test_an_unsigned_seed_pack_puts_nothing_into_gold(tmp_path):
    """The whole point. Every row marked `correct`, none of them signed."""
    path = _seed_pack(tmp_path, {}, sign=False)
    result = transcribe.read_workbook(path)
    assert result.records == {}, "unsigned rows must not reach the gold set"
    assert result.decisions
    assert all(d.destination == "review" for d in result.decisions)


def test_a_signed_seed_pack_reaches_gold(tmp_path):
    path = _seed_pack(tmp_path, {})
    result = transcribe.read_workbook(path)
    assert result.records, "signed, reviewed rows should reach the gold set"
    assert all(d.destination == "gold" for d in result.decisions)


def test_every_row_is_accounted_for_exactly_once(tmp_path):
    """No row may be silently dropped between the workbook and the two outputs."""
    path = _seed_pack(tmp_path, {})
    result = transcribe.read_workbook(path)
    gold_ids = {r["id"] for records in result.records.values() for r in records}
    decided = [d.record_id for d in result.decisions]
    assert len(decided) == len(set(decided)), "a row decided twice"
    assert gold_ids == {d.record_id for d in result.decisions if d.destination == "gold"}


def test_nothing_in_gold_is_ever_unsigned(tmp_path):
    """The invariant that must hold whatever the workbook contains."""
    path = _seed_pack(tmp_path, {"GC-001": "amend", "GE-001": "reject"})
    result = transcribe.read_workbook(path)
    for records in result.records.values():
        for record in records:
            assert record.get("approved_by") and record.get("approved_date")


def test_an_amend_verdict_holds_the_row_and_keeps_the_correction(tmp_path):
    path = _seed_pack(
        tmp_path,
        {"GC-001": "amend"},
        corrections={"GC-001": "the broader term is wrong"},
    )
    result = transcribe.read_workbook(path)
    held = {d.record_id: d for d in result.held}
    assert "GC-001" in held
    assert held["GC-001"].correction == "the broader term is wrong"
    assert "GC-001" not in {
        r["id"] for records in result.records.values() for r in records
    }


def test_an_unreadable_verdict_is_reported_and_the_row_is_held(tmp_path):
    path = _seed_pack(tmp_path, {"GC-001": "corrrect"})
    result = transcribe.read_workbook(path)
    assert any("corrrect" in str(p) for p in result.problems)
    assert "GC-001" in {d.record_id for d in result.held}


def test_a_plain_intake_workbook_still_transcribes_as_it_always_did(tmp_path):
    """No review columns means no gate — an authored workbook is not a seed pack."""
    gold = goldset.load(SOUND)
    records = {rt: list(rs) for rt, rs in gold.records.items()}
    book = workbook.build(generated="2026-08-25")
    workbook.fill(book, records)
    path = tmp_path / "plain.xlsx"
    book.save(path)

    result = transcribe.read_workbook(path)
    assert result.records, "an authored workbook transcribes without a verdict"
    assert result.decisions == []
