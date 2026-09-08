"""The concept typing pass — the 52, laid out for one person to sort.

The owner ruled on OQ-0001: *"Use those four groups — come back to me with the
list of 52 to sort."* This is the coming back.

**This module used to type nothing, and that changed on 2026-09-08.** Until
ADR-0079 every `type` cell was empty because a machine-filled taxonomy reads as
authoritative while having been authored by nobody. The owner replaced that
prohibition with a prohibition on *laundering*: an agent now types the concepts
and stamps every judgement `unreviewed`, with its model, its evidence and its
reasoning. So the pass reads `authored/concept-types.yaml` and pre-fills the
`type` column from it (ADR-0092).

**What did not change is whose answer counts.** A pre-filled cell is a proposal
to correct, never an answer already given. `approved_by` stays blank in every
row this module writes, the sheet says on its face that every value in it was
written by a machine and read by nobody, and the only thing that moves a typing
into `eval/gold/` is still `tmk-transcribe` reading a workbook in which a person
wrote a verdict **and** their name (ADR-0048, ADR-0086). Silence promotes
nothing: a reviewer who works past a row without changing it has signed nothing
and the record stays `unreviewed`.

Two artefacts, because they are for two different moments:

- `data/derived/concept-typing.xlsx` — the pass itself. It is an ordinary intake
  workbook with the `concept-types` sheet pre-filled, so `tmk-transcribe` reads
  it back with no new code: the single door into `eval/gold/` stays single
  (ADR-0048).
- `data/derived/reports/concept-typing.md` — the evidence, for reading on the
  dashboard or beside the spreadsheet. Labels, near-misses, the passages each
  concept was drawn from, and — since ADR-0092 — the machine's proposed group
  with the reasoning behind it, the readings it rejected and the thing it most
  expects to have got wrong. That last part is the point: a reviewer correcting
  a stated argument is doing something much cheaper than composing one.

`workbook.fill()` has existed since P7 and was deliberately not a command:
*"generating an empty workbook and pre-filling one with content are different
decisions, and only the first has been made."* The owner has now made the second
(ADR-0071).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from tm_knowledge.authored import store as authored_store
from tm_knowledge.config import REPO_ROOT
from tm_knowledge.stage0 import goldset
from tm_knowledge.upstream.loader import load_corpus

__all__ = [
    "WORKBOOK_PATH",
    "REPORT_PATH",
    "summarise",
    "rows",
    "write_workbook",
    "render",
    "write_report",
]

WORKBOOK_PATH = REPO_ROOT / "data" / "derived" / "concept-typing.xlsx"
REPORT_PATH = REPO_ROOT / "data" / "derived" / "reports" / "concept-typing.md"

#: What the four groups mean, in the owner's own question's words. Repeated on
#: the sheet and in the report so nobody has to hold them in their head across
#: 52 rows.
GROUPS: tuple[tuple[str, str], ...] = (
    ("ground_of_refusal", "a reason an application can be refused"),
    ("legal_test", "a question the decision maker has to answer"),
    ("relevant_factor", "something that feeds into that answer"),
    ("exception", "something that takes a case out of the rule"),
    ("none_of_these", "none of the four fit — which is an answer, not a gap"),
)


def summarise(concept: dict[str, Any]) -> str:
    """The concept in one line: what it is called, what else it is called, and
    what it is explicitly not.

    This is the `notes` cell of a typing row, and it is the whole of what a
    sorter sees without opening the evidence pack. All three parts come
    verbatim from the approved concept record — nothing here paraphrases or
    infers, because a summary that reworded a label would be a machine making a
    vocabulary judgement (CLAUDE.md rule 1).

    The alternative labels earn their place for the same reason the near-misses
    do. `pref_label` is one form of words out of several the Manual uses, so a
    sorter who does not recognise *connotation* may well recognise *secondary
    or implied meaning* — and the near-misses then stop them typing by label
    once they do.
    """
    summary = concept["pref_label"]
    if concept.get("alt_labels"):
        summary += " — also called: " + ", ".join(concept["alt_labels"])
    if concept.get("not_labels"):
        summary += " — not: " + ", ".join(concept["not_labels"])
    return summary


def _existing_types() -> dict[str, dict[str, Any]]:
    """Typings already signed, keyed by concept id. A second pass must not ask
    again about a concept somebody has already ruled on."""
    return {record["concept"]: record for record in goldset.load()["concept_type"]}


def _proposed_types(
    authored: authored_store.AuthoredSet | None = None,
) -> dict[str, authored_store.AuthoredRecord]:
    """Typings a machine proposed, keyed by concept id.

    Sound entries only. A refused authored record has no usable provenance, and
    a proposal whose envelope will not validate must not reach a reviewer's
    sheet looking exactly like one that would — that is the laundering ADR-0079
    prohibits, arriving through the back door of a spreadsheet.
    """
    authored = authored if authored is not None else authored_store.load()
    return {
        str(entry.record["concept"]): entry
        for entry in authored.of("concept_type")
        if entry.sound and isinstance(entry.record.get("concept"), str)
    }


def rows(
    gold: goldset.GoldSet | None = None,
    authored: authored_store.AuthoredSet | None = None,
) -> tuple[dict[str, Any], ...]:
    """One `concept_type` record per approved concept, in the order they rank.

    Three states a row can be in, and the difference between the second and the
    third is the whole of ADR-0092:

    - **signed** — a person has ruled on this concept. Carried through
      unchanged, and never re-asked.
    - **proposed** — `authored/` holds a machine's typing. The `type` cell is
      pre-filled from it and `approved_by` is left blank, so the reviewer is
      correcting a stated answer rather than composing one from nothing.
    - **neither** — `type` is absent, as every row was before ADR-0079.

    Ids are allocated by appending, never filling a gap left by a withdrawal
    (`IDENTIFIERS.md` §3). A concept already typed — signed *or* authored —
    keeps the id that typing carries: one `GT-0007` exists in this project and
    it moves between stores rather than being minted twice (ADR-0080).
    """
    gold = gold or goldset.load()
    proposed = _proposed_types(authored)
    concepts = sorted(gold["gold_concept"], key=lambda record: record["id"])
    already = {record["concept"]: record for record in gold["concept_type"]}

    taken = {
        int(record["id"].split("-")[1])
        for record in gold["concept_type"]
        if record.get("id", "").startswith("GT-")
    }
    # The authored store's ids are taken too. Without this the next pass mints
    # `GT-0001` a second time for a different concept, and two records claiming
    # one id across two stores is a defect the harness reports and nothing here
    # could undo (ADR-0080 consequence 1).
    taken |= {
        int(str(entry.record["id"]).split("-")[1])
        for entry in proposed.values()
        if str(entry.record.get("id", "")).startswith("GT-")
    }
    taken |= {
        int(identifier.split("-")[1])
        for identifier in gold.retired_ids
        if identifier.startswith("GT-")
    }
    next_number = max(taken, default=0) + 1

    built: list[dict[str, Any]] = []
    for concept in concepts:
        signed = already.get(concept["id"])
        if signed is not None:
            built.append(dict(signed))
            continue
        machine = proposed.get(concept["id"])
        if machine is not None:
            row = dict(machine.record)
            # The reviewer signs; the machine never does. Whatever the authored
            # record says about itself, these two cells leave here empty.
            row["approved_by"] = None
            row["approved_date"] = None
            row["notes"] = summarise(concept)
            built.append(row)
            continue
        built.append(
            {
                "id": f"GT-{next_number:04d}",
                "concept": concept["id"],
                # `type` is absent, not blank-stringed. A blank cell is a gap the
                # transcriber reports; an empty string would be a value.
                "notes": summarise(concept),
            }
        )
        next_number += 1
    return tuple(built)


def write_workbook(path: Path | None = None, *, generated: str | None = None) -> Path:
    """The pass, as a workbook `tmk-transcribe` already knows how to read."""
    from tm_knowledge.stage0 import workbook as workbook_module

    path = path or WORKBOOK_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    typing_rows = rows()
    proposed = sum(
        1
        for row in typing_rows
        if row.get("type") is not None and not row.get("approved_by")
    )
    book = workbook_module.build(generated=generated)
    _mark_as_scoped(book, len(typing_rows), proposed)
    workbook_module.fill(book, {"concept_type": list(typing_rows)})
    _fit_the_notes(book, len(typing_rows))
    book.save(path)
    return path


def _fit_the_notes(book, count: int) -> None:
    """Make the `notes` cell readable at the length it now is.

    The summary carries the concept's other names as well as its near-misses,
    so it is several times longer than a hand-typed note. At the intake
    workbook's default 28 characters it shows as one clipped line, and what
    falls off the end is the *not:* part — the half that stops a sorter typing
    by label. Widen the column and let it wrap; the rows auto-fit from there.

    Only the pre-filled rows are touched. The columns below them stay as the
    generated workbook made them, because this sheet is still an ordinary
    intake sheet that `tmk-transcribe` reads back (ADR-0048).
    """
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    from tm_knowledge.stage0.intake import sheets

    spec = next(sheet for sheet in sheets() if sheet.record_type == "concept_type")
    column = [item.header for item in spec.columns].index("notes") + 1
    sheet = book[spec.name]
    sheet.column_dimensions[get_column_letter(column)].width = 72
    for row in range(2, count + 2):
        sheet.cell(row=row, column=column).alignment = Alignment(
            wrap_text=True, vertical="top"
        )


def _mark_as_scoped(book, count: int, proposed: int = 0) -> None:
    """Say on the workbook's face what it is and is not.

    ADR-0055's third guard: a scoped artefact that looks exactly like an
    unscoped one is how a partial round gets filed as a complete one. This one
    is narrower still — every sheet but `concept-types` is empty and is meant to
    stay empty.

    Since ADR-0092 it carries a second warning, and it is the more important
    one: a sheet arriving with the answers already in it looks exactly like a
    sheet somebody else already worked through. The banner says how many values
    a machine wrote, that no person has read any of them, and where the
    reasoning is — because a pre-filled cell that does not announce itself is
    the laundering rule 1 prohibits, in a spreadsheet.
    """
    from openpyxl.styles import Alignment, Font

    sheet = book[book.sheetnames[0]]
    sheet["A1"] = "Concept typing pass — sort the concepts into the four groups"
    sheet["A1"].font = Font(bold=True, size=14)
    machine_note = (
        (
            f"**{proposed} of the {count} `type` cells are already filled in, and every one "
            "of them was written by a machine that no trade marks expert has checked.** "
            "They are proposals to correct, not answers. Nothing in this sheet is approved "
            "knowledge and nothing becomes approved by your leaving it alone — a row you "
            "read and do not change stays unreviewed. Change the ones that are wrong, and "
            "sign the ones you agree with; only a row carrying your name in `approved_by` "
            "counts as decided.\n\n"
            "The reasoning behind each proposal — why that group and not the obvious "
            "alternative, and the thing it most expects to have got wrong — is in "
            "`data/derived/reports/concept-typing.md`. Read a row's argument there before "
            "signing it.\n\n"
        )
        if proposed
        else ""
    )
    sheet["A2"] = (
        machine_note
        + f"This is not the general intake workbook. Exactly one sheet is in use — "
        f"**concept-types**, pre-filled with {count} rows, one per approved concept. Pick a "
        "group from the dropdown in the `type` column, put your name in `approved_by` and "
        "the date in `approved_date`, and hand the file back. Leave a row blank if you are "
        "not sure: a blank is reported as still-to-do, and a guess is not. Every other "
        "sheet is empty and should stay empty.\n\n"
        "The `notes` column is filled in for you and is the only column you should not "
        "need to touch: it gives each concept's name, what else the Manual calls it "
        "(*also called*), and what it is explicitly *not*. All three come straight from "
        "the approved concept record. For the passages behind them, read "
        "`data/derived/reports/concept-typing.md` beside this sheet."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 260 if proposed else 130
    row = 4
    for value, meaning in GROUPS:
        sheet.cell(row=row, column=1, value=value).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=meaning).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        row += 1


def _proposal_block(entry: authored_store.AuthoredRecord) -> list[str]:
    """One machine typing, laid out to be argued with rather than skimmed.

    The order is deliberate and it is not the order the envelope stores them in.
    *Why* comes first because it is the claim; *instead of* second because it is
    where an expert who already disagrees will look; *check* third because it is
    the concession, and a concession read before the argument reads as a
    disclaimer. The confidence comes last and is labelled as the machine's own —
    a number early in a block is read as authority, and this one is not evidence
    of anything (PU-0003).
    """
    envelope = entry.envelope or {}
    block: list[str] = [
        f"**Proposed group: `{entry.record.get('type', '—')}` — "
        f"written by `{entry.authored_by or 'an unnamed model'}`, "
        f"{entry.review_status}, read by no expert.**",
        "",
    ]
    reasoning = envelope.get("reasoning")
    if reasoning:
        block += [f"*Why:* {reasoning}", ""]
    alternatives = envelope.get("alternatives_considered") or ()
    if alternatives:
        block.append("*Instead of:*")
        block += [f"- {alternative}" for alternative in alternatives]
        block.append("")
    check = envelope.get("expert_should_check")
    if check:
        block += [f"*Check this first:* {check}", ""]
    confidence = envelope.get("confidence")
    basis = envelope.get("authoring_basis", "unknown")
    if confidence is not None:
        block += [
            f"*The machine's own confidence:* {confidence} · *basis:* `{basis}` — "
            "neither is evidence that the typing is right, and neither may be served to "
            "an examiner as one (PU-0003).",
            "",
        ]
    block += ["*The passages it was drawn from:*", ""]
    return block


def render(generated: str | None = None) -> str:
    """The evidence pack: every concept, what it is, what it is not, and what a
    machine says it is.

    The report was written to help somebody sort 52 concepts from a blank sheet.
    Since ADR-0092 it has a second and larger job: to let them disagree with an
    answer already given, cheaply and specifically. So each concept now carries
    the proposed group, the argument for it, the readings rejected and the thing
    the proposal most expects to have got wrong — and it says on every one of
    them that no expert has read a word of it.
    """
    gold = goldset.load()
    corpus = load_corpus()
    concepts = sorted(gold["gold_concept"], key=lambda record: record["id"])
    signed = _existing_types()
    proposed = _proposed_types()
    stamp = generated or date.today().isoformat()

    untyped = [c for c in concepts if c["id"] not in signed]
    machine_typed = [c for c in untyped if c["id"] in proposed]
    tally: dict[str, int] = {}
    for concept in machine_typed:
        group = str(proposed[concept["id"]].record.get("type", "—"))
        tally[group] = tally.get(group, 0) + 1

    parts: list[str] = [
        "<!-- Generated by tm_knowledge.stage0.typing. Do not hand-edit. -->",
        "",
        "# The concept typing pass — the evidence",
        "",
        f"**Generated {stamp}** by `tmk-typing --write`, from `eval/gold/concepts.yaml`, "
        "`authored/concept-types.yaml` and the pinned snapshot.",
        "",
        "## What this is",
        "",
        "You ruled on OQ-0001: *“Use those four groups — come back to me with the list of "
        f"52 to sort.”* Here is the list. There are **{len(concepts)}** approved concepts, "
        f"of which **{len(untyped)}** have not been sorted by a person.",
        "",
        "Sort them in `data/derived/concept-typing.xlsx`. The `type` column is a dropdown "
        "with the five values below; this document is the evidence to sort by, so keep it "
        "open beside the spreadsheet.",
        "",
        "| group | what it means |",
        "|---|---|",
        *[f"| `{value}` | {meaning} |" for value, meaning in GROUPS],
        "",
        "**Leaving a row blank is fine and is not the same as `none_of_these`.** A blank "
        "says *not yet sorted* and comes back on the next pass. `none_of_these` says *the "
        "four groups do not fit this one*, which is evidence about the taxonomy — and if a "
        "lot of rows come back that way, the taxonomy is what needs revisiting, not the "
        "rows.",
        "",
    ]

    if machine_typed:
        parts += [
            "## Every group below was proposed by a machine",
            "",
            f"**{len(machine_typed)} of the {len(untyped)} unsigned concepts already carry a "
            "proposed group, and no trade marks expert has read any of them.** They were "
            "written under ADR-0079, which replaced the rule that an agent may not author "
            "legal content with a rule that it may not pass authored content off as "
            "reviewed. Every proposal is stamped `unreviewed` in "
            "`authored/concept-types.yaml`, with the model that wrote it, the passages it "
            "rests on and the reasoning below.",
            "",
            "**Your job here is to disagree, specifically.** A proposal you leave alone is "
            "not thereby right and does not become approved: there are three states — "
            "unreviewed, approved, rejected — and only a row carrying your name in "
            "`approved_by` has moved (ADR-0086). Reading this document and changing nothing "
            "leaves all "
            f"{len(machine_typed)} exactly as unvalidated as they were before you opened "
            "it.",
            "",
            "Each concept below carries four things worth more than the group itself: "
            "**why** that group, **instead of** what, **check** — the thing the proposal "
            "most expects to have got wrong — and a confidence between 0 and 1 that is the "
            "machine's own and means nothing about whether it is right.",
            "",
            "| proposed group | concepts |",
            "|---|---|",
            *[
                f"| `{group}` | {tally[group]} |"
                for group, _ in GROUPS
                if group in tally
            ],
            "",
            "**The shape of that table is itself a finding.** The vocabulary was built "
            "around one ground of refusal — section 43 — so almost everything in it is "
            "material feeding that ground's question rather than a sibling ground. If "
            "`ground_of_refusal` looks too empty to you, the disagreement is about the "
            "taxonomy rather than about any single row, and it is worth saying so.",
            "",
        ]

    parts += ["## The concepts", ""]

    for concept in concepts:
        existing = signed.get(concept["id"])
        machine = proposed.get(concept["id"])
        heading = f"### `{concept['id']}` — {concept['pref_label']}"
        if existing:
            heading += (
                f"  *(already sorted: `{existing['type']}`, "
                f"{existing.get('approved_by') or 'unsigned'})*"
            )
        elif machine is not None:
            heading += (
                f"  *(machine proposes: `{machine.record.get('type', '—')}` — "
                "unreviewed, nobody has checked this)*"
            )
        parts.append(heading)
        parts.append("")
        if concept.get("alt_labels"):
            parts.append("*Also called:* " + ", ".join(concept["alt_labels"]))
        if concept.get("not_labels"):
            parts.append(
                "*Explicitly **not** the same as:* "
                + ", ".join(f"**{label}**" for label in concept["not_labels"])
            )
        if concept.get("broader") or concept.get("narrower"):
            wider = ", ".join(concept.get("broader") or ()) or "—"
            narrower = ", ".join(concept.get("narrower") or ()) or "—"
            parts.append(f"*Broader:* {wider} · *Narrower:* {narrower}")
        if concept.get("legislative_basis"):
            parts.append("*Legislative basis:* " + ", ".join(concept["legislative_basis"]))
        if concept.get("notes"):
            parts.append("")
            parts.append(f"> {concept['notes']}")
        parts.append("")
        if existing is None and machine is not None:
            parts += _proposal_block(machine)
        for ref in concept.get("definition_sources") or ():
            chunk = corpus.chunks.get(ref.split("~")[0]) or corpus.chunks.get(ref)
            if chunk is None:
                parts.append(f"- `{ref}` — not resolvable in the pinned snapshot")
                continue
            text = chunk.text
            excerpt = text if len(text) <= 700 else text[:700].rsplit(" ", 1)[0] + " …"
            parts.append(f"- **`{ref}`** — {excerpt}")
        parts.append("")

    return "\n".join(parts).rstrip() + "\n"


def write_report(path: Path | None = None, generated: str | None = None) -> Path:
    path = path or REPORT_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render(generated), encoding="utf-8")
    return path
